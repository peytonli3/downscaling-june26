#!/usr/bin/env python3
"""Standalone 1-slide .pptx explaining the CURRENT (sign-aware, per-tail-routed) extreme
weighting mechanism for the pinball loss -- see extreme_pixel_weights_signed in
train_new_enscgp_swin.py. Same visual style as build_presentation.py's main deck (shared
palette/helpers, duplicated here to keep this a self-contained single-purpose script).

This mechanism replaced an earlier design (shown in the "Loss III" slide of the original
deck) that ranked a single combined sqrt(u^2+v^2) magnitude and reused the identical weight
tensor for both q90 and q10 -- with no sign-awareness, so it pushed both tails equally hard
at every extreme pixel regardless of which direction the extreme actually was in. This slide
documents the fix.

Regenerate with:
    python extreme_weighting_slide.py
Output: runs/0714/presentation/extreme_weighting.pptx -- this script is source and lives
under extra_scripts/; its .pptx is a run artifact and belongs with the run.
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt

REPO = next(p for p in Path(__file__).resolve().parents
            if (p / "scripts" / "paths.py").is_file())
sys.path.insert(0, str(REPO / "scripts"))

from paths import RUNS_DIR  # noqa: E402

OUT = RUNS_DIR / "0714" / "presentation" / "extreme_weighting.pptx"

# ── palette (matches build_presentation.py) ──
NAVY = RGBColor(0x14, 0x2A, 0x4A)
ACCENT = RGBColor(0x2E, 0x86, 0xAB)
GOOD = RGBColor(0x2E, 0x7D, 0x46)
BAD = RGBColor(0xB3, 0x3A, 0x3A)
GRAY = RGBColor(0x33, 0x33, 0x33)
LIGHT = RGBColor(0xEC, 0xF2, 0xF6)
EQBG = RGBColor(0xF4, 0xF1, 0xE8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SW, SH = 13.333, 7.5

prs = Presentation()
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH)
BLANK = prs.slide_layouts[6]


def _txt(slide, l, t, w, h, anchor=None):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    if anchor is not None:
        tf.vertical_anchor = anchor
    return tf


def _rect(slide, l, t, w, h, color):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    sp.fill.solid()
    sp.fill.fore_color.rgb = color
    sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def _set(p, text, size, color=GRAY, bold=False, align=None, font="Calibri"):
    p.text = text if text else " "
    r = p.runs[0]
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = font
    if align is not None:
        p.alignment = align


def header(slide, title):
    tf = _txt(slide, 0.6, 0.32, SW - 1.2, 0.95, anchor=MSO_ANCHOR.MIDDLE)
    _set(tf.paragraphs[0], title, 28, NAVY, bold=True)
    _rect(slide, 0.62, 1.24, SW - 1.24, 0.045, ACCENT)


def equation(slide, l, t, w, lines, size=17):
    h = 0.4 * len(lines) + 0.34
    _rect(slide, l, t, w, h, EQBG)
    _rect(slide, l, t, 0.08, h, ACCENT)
    tf = _txt(slide, l + 0.28, t + 0.14, w - 0.5, h - 0.24)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        _set(p, line, size, NAVY, bold=False, font="Consolas")
        p.space_after = Pt(3)
    return t + h


slide = prs.slides.add_slide(BLANK)
header(slide, "Extreme weighting, sign-aware (current)")

bot = equation(slide, 0.6, 1.5, SW - 1.2, [
    "w = 1 + α·F(|y|)         F = per-(sample, component) empirical CDF of |target|",
    "w ← w / mean(w)            detached — no gradient through the ranking",
    "",
    "w_upper = w  if y ≥ 0  else 1         w_lower = w  if y < 0  else 1",
])

tf = _txt(slide, 0.6, bot + 0.22, SW - 1.2, SH - bot - 1.0)
items = [
    ("Judged by magnitude, ROUTED by sign", 0, ACCENT),
    ("u and v are signed — a strong westward gust is exactly as extreme as an equally "
     "strong eastward one — so severity is ranked from |target| (sign-blind), then each "
     "pixel's weight is routed to exactly ONE tail by its sign.", 1, GRAY),
    ("Ranked per component, not combined", 0, GRAY),
    ("u and v are each ranked against their OWN distribution (not a shared √(u²+v²) "
     "magnitude) — a pixel extreme only in u doesn't spuriously also weight v.", 1, GRAY),
    ("The tail that doesn't apply gets baseline weight 1", 0, GRAY),
    ("No wasted pressure: at a negative-u pixel, q90 (nothing to reach for there) isn't "
     "boosted — only q10 is.", 1, GRAY),
    ("Why not rank the two signs separately?", 0, BAD),
    ("Tried first — wrong: on a one-directional sample (severe negative gusts, only mild "
     "positive values), independently-normalized ranks would boost the \"most positive\" "
     "pixel toward the SAME max weight as the true severe extreme, despite not being severe "
     "in absolute terms. Ranking |target| once fixes this.", 1, BAD),
]
first = True
for text, level, color in items:
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    first = False
    prefix = "▸  " if level == 0 else "     "
    _set(p, prefix + text, 16 if level == 0 else 14, color, bold=(level == 0))
    p.space_after = Pt(8 if level == 0 else 12)

_rect(slide, 0.6, SH - 0.85, SW - 1.2, 0.58, LIGHT)
_rect(slide, 0.6, SH - 0.85, 0.09, 0.58, ACCENT)
tf2 = _txt(slide, 0.85, SH - 0.83, SW - 1.5, 0.54, anchor=MSO_ANCHOR.MIDDLE)
_set(tf2.paragraphs[0],
     "Replaces a design that reused one shared magnitude weight identically for both tails "
     "(w10 = w90) — which pushed both equally hard regardless of which direction the extreme was in.",
     13, NAVY, bold=True)

prs.save(str(OUT))
print(f"Saved {OUT}")