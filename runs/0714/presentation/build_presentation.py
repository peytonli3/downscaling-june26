#!/usr/bin/env python3
"""Build the q10/q50/q90 model presentation (v0714) as a .pptx, pulling the generated
figures from runs/0714/figures/. Audience: the lab (assumes familiarity with the
ERA5/EnsCGP/Swin pipeline) -- minimal intro, deep on the loss and the results.
Regenerate with:
    python build_presentation.py
Output: q10_q50_q90_model.pptx next to this script. Speaker notes (with rough timing) are
attached to every slide."""
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
FIG = HERE.parent / "figures"
OUT = HERE / "q10_q50_q90_model.pptx"

# ── palette ──
NAVY = RGBColor(0x14, 0x2A, 0x4A)
ACCENT = RGBColor(0x2E, 0x86, 0xAB)
GOOD = RGBColor(0x2E, 0x7D, 0x46)
BAD = RGBColor(0xB3, 0x3A, 0x3A)
GRAY = RGBColor(0x33, 0x33, 0x33)
LIGHT = RGBColor(0xEC, 0xF2, 0xF6)
EQBG = RGBColor(0xF4, 0xF1, 0xE8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SW, SH = 13.333, 7.5

prs = Presentation()
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH)
BLANK = prs.slide_layouts[6]


def _txt(slide, l, t, w, h, anchor=None):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    if anchor is not None:
        tf.vertical_anchor = anchor
    return tf


def _rect(slide, l, t, w, h, color):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    sp.fill.solid()
    sp.fill.fore_color.rgb = color
    sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def _set(p, text, size, color=GRAY, bold=False, align=None, font="Calibri"):
    # A truly empty string yields a paragraph with no runs, so blank spacer lines use a
    # space instead (keeps the run available for font styling).
    p.text = text if text else " "
    r = p.runs[0]
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = font
    if align is not None:
        p.alignment = align


def _footer(slide, n):
    tf = _txt(slide, SW - 1.2, SH - 0.5, 0.9, 0.35)
    _set(tf.paragraphs[0], f"{n}", 12, RGBColor(0x99, 0x99, 0x99), align=PP_ALIGN.RIGHT)


def header(slide, title):
    tf = _txt(slide, 0.6, 0.32, SW - 1.2, 0.95, anchor=MSO_ANCHOR.MIDDLE)
    _set(tf.paragraphs[0], title, 30, NAVY, bold=True)
    _rect(slide, 0.62, 1.28, SW - 1.24, 0.045, ACCENT)


def equation(slide, l, t, w, lines, size=18):
    """Monospace equation block on a tinted background."""
    h = 0.42 * len(lines) + 0.34
    _rect(slide, l, t, w, h, EQBG)
    _rect(slide, l, t, 0.08, h, ACCENT)
    tf = _txt(slide, l + 0.28, t + 0.14, w - 0.5, h - 0.24)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        _set(p, line, size, NAVY, bold=False, font="Consolas")
        p.space_after = Pt(4)
    return t + h


def bullets(n, title, items, notes="", top=1.55, size0=21, size1=17):
    slide = prs.slides.add_slide(BLANK)
    header(slide, title)
    tf = _txt(slide, 0.75, top, SW - 1.5, SH - top - 0.6)
    first = True
    for item in items:
        text, level = item[0], item[1]
        color = item[2] if len(item) > 2 else GRAY
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        _set(p, ("▸  " if level == 0 else "–  ") + text, size0 if level == 0 else size1,
             color, bold=(level == 0))
        p.level = level
        p.space_after = Pt(9 if level == 0 else 4)
        if level == 1:
            p.runs[0].font.bold = False
    _footer(slide, n)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def fit(img_path, bl, bt, bw, bh):
    iw, ih = Image.open(img_path).size
    ar, box_ar = iw / ih, bw / bh
    w, h = (bw, bw / ar) if ar > box_ar else (bh * ar, bh)
    return bl + (bw - w) / 2, bt + (bh - h) / 2, w, h


def figure(n, title, img, takeaway, notes=""):
    slide = prs.slides.add_slide(BLANK)
    header(slide, title)
    l, t, w, h = fit(img, 0.5, 1.5, SW - 1.0, 5.0)
    slide.shapes.add_picture(str(img), Inches(l), Inches(t), Inches(w), Inches(h))
    _rect(slide, 0.6, SH - 0.92, SW - 1.2, 0.62, LIGHT)
    _rect(slide, 0.6, SH - 0.92, 0.09, 0.62, ACCENT)
    tf = _txt(slide, 0.85, SH - 0.9, SW - 1.5, 0.58, anchor=MSO_ANCHOR.MIDDLE)
    _set(tf.paragraphs[0], takeaway, 15, NAVY, bold=True)
    _footer(slide, n)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


# ── 1. TITLE ──
s = prs.slides.add_slide(BLANK)
_rect(s, 0, 0, SW, SH, NAVY)
_rect(s, 0, 4.35, SW, 0.06, ACCENT)
tf = _txt(s, 0.9, 2.2, SW - 1.8, 2.0, anchor=MSO_ANCHOR.BOTTOM)
_set(tf.paragraphs[0], "Direct Quantile Regression", 44, WHITE, bold=True)
p = tf.add_paragraph()
_set(p, "Replacing the Gaussian uncertainty head  ·  q10 / q50 / q90", 25, RGBColor(0xBF, 0xD7, 0xE6))
tf2 = _txt(s, 0.92, 4.6, SW - 1.8, 1.4)
_set(tf2.paragraphs[0], "Loss design and calibration results for the EnsCGP → Swin2SR refiner", 18,
     RGBColor(0xD8, 0xE4, 0xEC))
p = tf2.add_paragraph()
_set(p, "model version 0714  ·  git tag v6-0714  ·  July 2026", 15, RGBColor(0x9F, 0xB8, 0xC8))
s.notes_slide.notes_text_frame.text = (
    "Lab audience — skip the pipeline recap, they know it. Frame: this is about the loss "
    "design and what the calibration diagnostics actually show. [~0:20]")

# ── 2. WHAT CHANGED (compressed intro) ──
sl = bullets(2, "What changed in v0714", [
    ("Pipeline unchanged: ERA5 → EnsCGP analog posterior → Swin2SR refiner (2.09M params)", 0),
    ("Backbone, terrain encoder, structural losses, data pipeline all identical", 1),
    ("Old head: μ + 2×2 Cholesky covariance, trained by Gaussian NLL", 0, BAD),
    ("Symmetric and light-tailed; NLL ties the spread to a shape we don't believe at peaks", 1),
    ("New head: direct q10 / q50 / q90 per component (u, v) — distribution-free", 0, GOOD),
    ("Only the OUTPUT HEAD and the UNCERTAINTY LOSS changed → results are attributable", 1),
], notes=(
    "One slide of context. The key framing for the lab: this is a surgical change, so any "
    "difference in results is attributable to the head + loss, not confounded by backbone or "
    "data changes. [~1:00]"), top=1.7)

# ── 3. HEAD: MONOTONE BY CONSTRUCTION ──
s = prs.slides.add_slide(BLANK)
header(s, "The head — monotone by construction")
bot = equation(s, 0.75, 1.55, SW - 1.5, [
    "up   = softplus( inv_softplus(1.28·σ_EnsCGP) + g·raw_up   ) + ε        > 0",
    "down = softplus( inv_softplus(1.28·σ_EnsCGP) + g·raw_down ) + ε        > 0",
    "q90 = q50 + up        q10 = q50 − down       ⇒   q10 ≤ q50 ≤ q90",
], size=16)
tf = _txt(s, 0.75, bot + 0.28, SW - 1.5, SH - bot - 1.0)
for i, (t_, lv, c) in enumerate([
    ("Quantile crossing is impossible — guaranteed, not penalized", 0, GOOD),
    ("Offsets are strictly positive, so no crossing penalty term is needed at all", 1, GRAY),
    ("Offsets seeded from the EnsCGP per-pixel σ  (σ_u = L11, σ_v = √(L21²+L22²))", 0, GRAY),
    ("A fresh model emits q50 ± 1.28σ — the EnsCGP posterior re-expressed as an 80% band", 1, GRAY),
    ("softplus, never ReLU", 0, GRAY),
    ("ReLU's dead zone lets the spread collapse to exactly zero with zero gradient — unrecoverable", 1, GRAY),
]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    _set(p, ("▸  " if lv == 0 else "–  ") + t_, 20 if lv == 0 else 16, c, bold=(lv == 0))
    p.space_after = Pt(7)
_footer(s, 3)
s.notes_slide.notes_text_frame.text = (
    "Monotonicity is structural, not learned — worth stressing, since crossing penalties are the "
    "usual hack. The EnsCGP-σ seeding means we start from a physically sensible band and only "
    "learn corrections. The softplus/ReLU point is a real failure mode we designed around. [~1:30]")

# ── 4. LOSS I: PINBALL ──
s = prs.slides.add_slide(BLANK)
header(s, "Loss I — the pinball (quantile) loss")
bot = equation(s, 0.75, 1.55, SW - 1.5, [
    "L_τ(q, y) = max( τ·(y − q) ,  (τ − 1)·(y − q) )",
    "          = τ·(y−q)      if y > q        (under-predicted)",
    "            (1−τ)·(q−y)  if y ≤ q        (over-predicted)",
], size=17)
tf = _txt(s, 0.75, bot + 0.25, SW - 1.5, SH - bot - 1.0)
for i, (t_, lv, c) in enumerate([
    ("Asymmetric by design: at τ = 0.9, missing LOW costs 9× missing high", 0, GRAY),
    ("So the minimizer is pushed up until only 10% of truth lies above it", 1, GRAY),
    ("argmin over a sample = the empirical τ-quantile (we unit-test this recovery)", 0, ACCENT),
    ("Distribution-free: no shape assumption anywhere, unlike NLL", 1, GRAY),
    ("Ties directly to CRPS:   CRPS(F, y) = 2 ∫₀¹ L_τ(F⁻¹(τ), y) dτ", 0, GRAY),
    ("Pinball is the proper scoring rule per quantile; CRPS is its integral", 1, GRAY),
]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    _set(p, ("▸  " if lv == 0 else "–  ") + t_, 20 if lv == 0 else 16, c, bold=(lv == 0))
    p.space_after = Pt(7)
_footer(s, 4)
s.notes_slide.notes_text_frame.text = (
    "Spend real time here. The asymmetry IS the mechanism — walk through the τ=0.9 case out "
    "loud. Mention we verified the minimizer recovers the true quantile of a known distribution "
    "in a unit test, so the implementation is trustworthy. The CRPS link matters because it's "
    "how we score the model later. [~1:45]")

# ── 5. LOSS II: THE q50 DESIGN DECISION ──
bullets(5, "Loss II — why q50 is NOT pinball-trained", [
    ("Pinball at τ = 0.5 is just ½·L1 → drives q50 to the POINTWISE median", 0, BAD),
    ("Our dominant error mode is spatial DISPLACEMENT, not amplitude", 0),
    ("Pointwise metrics reward blurring: hedge across where the feature might be", 1),
    ("An earlier per-band diagnostic showed exactly this — L1/spectral rewarded smoothing", 1),
    ("So q50 is trained ONLY by the displacement-tolerant structural losses", 0, GOOD),
    ("Laplacian-pyramid sliced-Wasserstein + frequency-band L1 — unchanged from before", 1),
    ("Adding a pinball(0.5) term would actively fight them → deliberately omitted", 1),
    ("Consequence: q50 is a SHARP central field, not a calibrated median…", 0, ACCENT),
    ("…yet it still lands at ~0.48 / 0.53 empirical coverage. Near-median for free", 1),
], notes=(
    "This is the most interesting slide for this audience — a genuine design decision with a "
    "tradeoff, not a default. The punchline: we refused to pinball-train the median to protect "
    "sharpness, and it came out near-median-calibrated anyway. Expect questions here. [~2:00]"),
    top=1.6, size0=20, size1=16)

# ── 6. LOSS III: EXTREME WEIGHTING ──
s = prs.slides.add_slide(BLANK)
header(s, "Loss III — extreme weighting on the upper tail")
bot = equation(s, 0.75, 1.5, SW - 1.5, [
    "w = 1 + α · F̂(|y|)          F̂ = per-sample empirical CDF (rank/N) of |wind|",
    "w ← w / mean(w)             detached — no gradient through the ranking",
], size=16)
tf = _txt(s, 0.75, bot + 0.22, SW - 1.5, SH - bot - 1.0)
for i, (t_, lv, c) in enumerate([
    ("Problem: damaging peaks are a tiny fraction of pixels", 0, BAD),
    ("Aggregate pinball under-trains them → q90 gets smoothed DOWN toward neighbours", 1, GRAY),
    ("Calm pixels keep w ≈ 1; the windiest approach w ≈ 1 + α", 0, GRAY),
    ("Mean-normalized, so α changes WHERE emphasis goes, not the overall loss scale", 1, GRAY),
    ("Applied to the q90 pinball (the damaging tail); optional on q10", 0, GRAY),
    ("Verified in a controlled shared-capacity test: q90 offset 0.82 → 1.13 with it on", 0, GOOD),
]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    _set(p, ("▸  " if lv == 0 else "–  ") + t_, 20 if lv == 0 else 16, c, bold=(lv == 0))
    p.space_after = Pt(7)
_footer(s, 6)
s.notes_slide.notes_text_frame.text = (
    "α is THE knob for the tail problem shown later. Two implementation details worth stating: "
    "detached (the ranking is a fixed weighting, gradient flows only through the pinball error) "
    "and mean-normalized (so sweeping α doesn't secretly rescale the loss). [~1:30]")

# ── 7. FULL OBJECTIVE ──
s = prs.slides.add_slide(BLANK)
header(s, "The full objective")
bot = equation(s, 0.75, 1.55, SW - 1.5, [
    "L =  w_ms  · MultiScale(q50)          ← displacement-tolerant, sharp",
    "   + w_freq· FreqBand(q50)",
    "   + w_pin · [ L_0.9(q90) + L_0.1(q10) ]      ← pinball, extreme-weighted",
    "",
    "current:  w_ms = 1.0    w_freq = 0.5    w_pin = 1.0    α tunable",
], size=16)
tf = _txt(s, 0.75, bot + 0.3, SW - 1.5, SH - bot - 1.1)
for i, (t_, lv, c) in enumerate([
    ("Clean separation of duties — each term owns exactly one output", 0, ACCENT),
    ("Structural terms NEVER touch q10/q90", 0, GRAY),
    ("An uncertainty envelope is not a wind field; matching its spectrum is a category error", 1, GRAY),
    ("Pinball NEVER touches q50 (slide 5)", 0, GRAY),
    ("The Gaussian NLL term is gone entirely — no covariance is predicted any more", 0, GRAY),
]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    _set(p, ("▸  " if lv == 0 else "–  ") + t_, 19 if lv == 0 else 16, c, bold=(lv == 0))
    p.space_after = Pt(6)
_footer(s, 7)
s.notes_slide.notes_text_frame.text = (
    "Pull the three loss slides together. The 'category error' line is worth saying explicitly — "
    "it's a guardrail written into the code docstrings so nobody wires the structural loss onto "
    "the envelopes later. [~1:00]")

# ── 8-12. RESULTS ──
figure(8, "Training & calibration trajectory", FIG / "training_curves.png",
       "Val loss 5.78 → 4.83. Coverage locks onto nominal within a few epochs: q90 ≈ 0.90, q10 ≈ 0.09, and holds.",
       notes=(
        "Panel 1 val loss; panel 2 pinball split into q90/q10; panel 3 structural terms; panel 4 "
        "coverage. The story: coverage snaps to nominal early and stays flat — the pinball term "
        "does its job almost immediately. Dashed extreme-tail lines sit visibly lower, "
        "foreshadowing the problem. [~1:15]"))

figure(9, "Qualitative results (held-out test)", FIG / "swin_quantile_panels.png",
       "q50 recovers structure the EnsCGP first guess lacks; the band brackets truth; CRPS concentrates on high-wind regions.",
       notes=(
        "ERA5 LR → EnsCGP → WRF truth, then q10/q50/q90 and CRPS. Say out loud: the q10/q90 "
        "panels are component scenario fields, NOT speed percentiles — the head gives marginal "
        "u/v quantiles with no correlation, so speed(q10_u,q10_v) is not a lower speed bound. "
        "CRPS lights up exactly where the wind is strong. [~1:30]"))

figure(10, "Spectral fidelity of the central field", FIG / "eigenspectra_aggregate_swin.png",
       "q50 tracks WRF's power spectrum across all scales; bicubic and EnsCGP fall short below ~100 km.",
       notes=(
        "This is the payoff for the slide-5 decision: because we refused to pinball-train q50, it "
        "kept the multi-scale loss's sharpness and matches WRF's spectrum. Direct evidence the "
        "design choice worked. [~1:00]"))

figure(11, "Calibration: PIT / rank histogram", FIG / "quantile_pit_histogram.png",
       "Bulk is calibrated — bins match nominal [.10 .40 .40 .10]. Extreme tail is NOT: the ≥q90 bin reaches ~0.15 (u) / ~0.18 (v).",
       notes=(
        "Fraction of truth landing in each of the four bins cut by the quantiles. Blue = overall, "
        "sits on the nominal steps. Red = extreme tail: the ≥q90 bin is 1.5–1.8× too tall, i.e. "
        "truth punches through q90 far too often at peaks, and the inner bins are correspondingly "
        "depleted. This is the headline limitation. [~1:45]"))

figure(12, "Spatial structure of the error", FIG / "quantile_coverage_maps.png",
       "q10/q90 near nominal spatially — but the q50 maps show a coherent land/ocean pattern: a systematic median bias over terrain.",
       notes=(
        "Per-pixel coverage minus nominal; blue = under-covers. The unexpected find is the q50 "
        "column — a strong land/ocean structure means the median is systematically biased over "
        "terrain, completely hidden by the aggregate 0.48/0.53 numbers. A second, independent "
        "lead. Caveat: current figure is from a 96-sample subset; full test split pending. [~1:30]"))

# ── 13. FINDINGS ──
bullets(13, "Findings", [
    ("Monotone by construction — zero crossing violations, no penalty term", 0, GOOD),
    ("Calibrated in the bulk: q90 ≈ 0.90 / q10 ≈ 0.09, PIT bins on nominal", 0, GOOD),
    ("Central field stayed sharp: matches WRF's spectrum across scales", 0, GOOD),
    ("Tail overconfidence: truth exceeds q90 ~1.5–1.8× too often at peaks", 0, BAD),
    ("Systematic q50 bias over terrain — invisible to aggregate metrics", 0, BAD),
    ("No baseline skill comparison yet — CRPS vs EnsCGP/bicubic still to do", 0, BAD),
], notes=(
    "Honest scorecard. Flag the third limitation yourself before someone asks it: we have not "
    "yet shown the quantiles beat the EnsCGP first guess on a probabilistic score. [~1:00]"))

# ── 14. NEXT STEPS ──
bullets(14, "Next steps", [
    ("Raise extreme-weight α; re-run PIT + coverage to confirm the tail closes", 0),
    ("CRPS skill score vs EnsCGP and bicubic — quantify probabilistic value added", 0),
    ("Sharpness check: interval width vs EnsCGP ±1.28σ (are we widening or improving?)", 0),
    ("Coverage vs wind-magnitude curve — resolve exactly where the tail breaks down", 0),
    ("Terrain-stratified coverage / CRPS — test the orographic hypothesis for the q50 bias", 0),
], notes=(
    "Prioritized and concrete. The middle three are the ones that answer a skeptical lab: is it "
    "better than the first guess, are you just inflating the band, and precisely where does it "
    "fail. [~1:00]"))

# ── 15. SUMMARY ──
s = prs.slides.add_slide(BLANK)
_rect(s, 0, 0, SW, SH, NAVY)
tf = _txt(s, 0.9, 1.1, SW - 1.8, 1.0)
_set(tf.paragraphs[0], "Summary", 34, WHITE, bold=True)
_rect(s, 0.92, 2.05, 3.0, 0.05, ACCENT)
tf = _txt(s, 0.95, 2.4, SW - 1.9, 4.3)
for i, (text, lead) in enumerate([
    ("Gaussian/Cholesky head → monotone q10/q50/q90 with a pinball objective", True),
    ("Crossing impossible by construction; offsets seeded from the EnsCGP posterior spread", False),
    ("q50 deliberately excluded from the pinball term to protect sharpness — and it worked", True),
    ("Spectrum matches WRF; empirical median coverage ~0.5 anyway", False),
    ("Calibrated in the bulk, overconfident in the damaging tail", True),
    ("α tuning + baseline skill scores are the immediate next steps", False),
]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    _set(p, ("→  " if lead else "     •  ") + text, 20 if lead else 16,
         WHITE if lead else RGBColor(0xC9, 0xDA, 0xE6), bold=lead)
    p.space_after = Pt(11)
tf = _txt(s, 0.95, 6.75, SW - 1.9, 0.5)
_set(tf.paragraphs[0], "git tag v6-0714  ·  every figure regenerable from a committed script", 15,
     RGBColor(0x9F, 0xB8, 0xC8))
s.notes_slide.notes_text_frame.text = (
    "Three beats: what changed, the design decision that paid off, and the honest open problem. "
    "Then open for questions. [~0:45]")

n_slides = len(prs.slides._sldIdLst)
prs.save(str(OUT))
print(f"Saved {OUT}  ({n_slides} slides)")